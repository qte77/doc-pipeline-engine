# Copyright 2026 qte77
#
# Licensed under the Apache License, Version 2.0 (the "License");
# you may not use this file except in compliance with the License.
# You may obtain a copy of the License at
#
#     http://www.apache.org/licenses/LICENSE-2.0
"""Office adapter — plain ``.docx`` / ``.xlsx`` text via python-docx + openpyxl.

A lightweight alternative to Kreuzberg for native Office files: no Kreuzberg
dependency and no ELv2 terms (see ADR-0005). Opt-in extra:
``uv sync --extra office``.

- ``.docx``: paragraphs and table rows in document order; cells tab-separated.
- ``.xlsx``: one ``# <sheet>`` line per sheet, then non-empty rows
  tab-separated (cached cell values, not formulas).
"""
from __future__ import annotations

from datetime import UTC, datetime
from typing import TYPE_CHECKING, Any

from doc_pipeline_engine.base.adapter import AdapterBase, register
from doc_pipeline_engine.models.extraction_bundle import (
    AdapterInfo,
    Content,
    ExtractionBundle,
)

if TYPE_CHECKING:
    from pathlib import Path

ADAPTER_NAME = "office"
ADAPTER_VERSION = "0.1.0"
_INSTALL_HINT = "Install the extra: uv sync --extra office"


def _docx_text(path: Path) -> str:
    try:
        import docx
        from docx.table import Table
    except ImportError as exc:
        raise RuntimeError(f"python-docx not installed. {_INSTALL_HINT}") from exc

    lines: list[str] = []
    for block in docx.Document(str(path)).iter_inner_content():
        if isinstance(block, Table):
            lines.extend("\t".join(cell.text for cell in row.cells) for row in block.rows)
        elif block.text.strip():
            lines.append(block.text)
    return "\n".join(lines)


def _xlsx_text(path: Path) -> str:
    try:
        import openpyxl
    except ImportError as exc:
        raise RuntimeError(f"openpyxl not installed. {_INSTALL_HINT}") from exc

    wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    lines: list[str] = []
    try:
        for ws in wb.worksheets:
            lines.append(f"# {ws.title}")
            for row in ws.iter_rows(values_only=True):
                if any(v is not None for v in row):
                    lines.append("\t".join("" if v is None else str(v) for v in row))
    finally:
        wb.close()
    return "\n".join(lines)


_READERS = {".docx": _docx_text, ".xlsx": _xlsx_text}


class OfficeAdapter(AdapterBase):
    """Extraction backend for ``.docx`` / ``.xlsx`` using python-docx + openpyxl."""

    name = ADAPTER_NAME
    version = ADAPTER_VERSION

    def extract(self, manifest_file: dict[str, Any], source_root: Path) -> ExtractionBundle:
        """Read one ``.docx`` / ``.xlsx`` file into an ExtractionBundle.

        Raises:
            ValueError: if the file is not ``.docx`` or ``.xlsx``.
            RuntimeError: if the needed library is not installed.
        """
        abs_path = source_root / manifest_file["path"]
        reader = _READERS.get(abs_path.suffix.lower())
        if reader is None:
            raise ValueError(
                f"office adapter handles .docx/.xlsx only, got {abs_path.suffix!r}"
            )
        return ExtractionBundle(
            source_path=manifest_file["path"],
            source_sha256=manifest_file["sha256"],
            adapter=AdapterInfo(name=ADAPTER_NAME, version=ADAPTER_VERSION),
            extracted_at=datetime.now(UTC).isoformat(),
            content=Content(text=reader(abs_path), layout=[]),
        )


register(OfficeAdapter())
